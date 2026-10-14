import datetime as dt, os, tempfile
os.environ["DATABASE_URL"] = f"sqlite:///{tempfile.mkdtemp()}/t.db"; os.environ["CRON_SECRET"] = "cron-secret"
import pytest
from fastapi.testclient import TestClient
from festops.main import app
from festops.models import Base, Event, EventScope, EventSession, Issue, SessionLocal, User, engine, now
from festops.security import hash_pw

PW = "correct-horse-battery"
@pytest.fixture(scope="module")
def env():
    Base.metadata.create_all(engine); db = SessionLocal(); ids = {}
    for k, r in dict(core="CORE", dep="DEPUTY_CORE", sup="SUPER_COORDINATOR", co="COORDINATOR", co2="COORDINATOR", v1="VOLUNTEER", v2="VOLUNTEER").items():
        u = User(email=f"{k}@t.invalid", name=k, role=r, pw_hash=hash_pw(PW)); db.add(u); db.flush(); ids[k] = u.id
    e1, e2 = Event(code="E1", name="One", status="ACTIVE"), Event(code="E2", name="Two", status="ACTIVE"); db.add_all([e1, e2]); db.flush()
    t = now().replace(microsecond=0)
    s1 = EventSession(event_id=e1.id, stage="A", starts_at=t, ends_at=t + dt.timedelta(hours=1))
    s2 = EventSession(event_id=e2.id, stage="B", starts_at=t + dt.timedelta(minutes=70), ends_at=t + dt.timedelta(hours=2))  # inside turnaround window of s1
    db.add_all([s1, s2]); db.flush(); db.add(EventScope(user_id=ids["co"], event_id=e1.id)); db.commit()
    ids.update(e1=e1.id, e2=e2.id, s1=s1.id, s2=s2.id); return ids
c = TestClient(app)
def H(k): return {"Authorization": "Bearer " + c.post("/api/v1/auth/login", json={"email": f"{k}@t.invalid", "password": PW}).json()["access_token"]}

def test_auth_required_and_bad_login(env):
    assert c.get("/api/v1/events").status_code == 401
    assert c.post("/api/v1/auth/login", json={"email": "core@t.invalid", "password": "wrong"}).status_code == 401

def test_role_permissions(env):
    assert c.post("/api/v1/events", json={"code": "X", "name": "X"}, headers=H("v1")).status_code == 403
    assert c.post("/api/v1/events", json={"code": "X", "name": "X"}, headers=H("co")).status_code == 403
    assert c.post("/api/v1/events", json={"code": "X", "name": "X"}, headers=H("dep")).status_code == 201
    assert c.get("/api/v1/audit-logs", headers=H("v1")).status_code == 403
    assert c.get("/api/v1/audit-logs", headers=H("core")).status_code == 200

def test_deputy_cannot_create_core_or_self_promote(env):
    body = {"email": "n@t.invalid", "name": "N", "role": "CORE", "password": PW}
    assert c.post("/api/v1/users", json=body, headers=H("dep")).status_code == 403
    assert c.post(f"/api/v1/users/{env['dep']}/roles", json={"role": "CORE"}, headers=H("dep")).status_code == 403
    assert c.post(f"/api/v1/users/{env['core']}/roles", json={"role": "CORE"}, headers=H("core")).status_code == 403  # own role

def test_event_scope_isolation(env):
    assert c.get(f"/api/v1/events/{env['e1']}", headers=H("co")).status_code == 200
    assert c.get(f"/api/v1/events/{env['e2']}", headers=H("co")).status_code == 404
    assert c.get(f"/api/v1/events/{env['e1']}", headers=H("v1")).status_code == 404

def test_conflict_detection_and_override(env):
    d = H("dep"); a = {"session_id": env["s1"], "user_id": env["v1"], "kind": "POC"}
    r = c.post("/api/v1/duties", json=a, headers=d); assert r.status_code == 201
    b = {"session_id": env["s2"], "user_id": env["v1"], "kind": "POC"}
    assert c.post("/api/v1/duties", json=b, headers=d).status_code == 409
    assert c.post("/api/v1/duties", json=b, headers=H("co")).status_code == 403
    assert c.post("/api/v1/duties", json=b | {"override_reason": "Short staffed"}, headers=d).status_code == 201

def test_checkin_idempotent_and_owner_only(env):
    duty = c.get("/api/v1/duties/my-duties", headers=H("v1")).json()[0]["id"]
    assert c.post(f"/api/v1/duties/{duty}/check-in", headers=H("v2")).status_code == 403
    first = c.post(f"/api/v1/duties/{duty}/check-in", headers=H("v1")).json()
    again = c.post(f"/api/v1/duties/{duty}/check-in", headers=H("v1")).json()
    assert first["checked_in_at"] == again["checked_in_at"]

def test_issue_workflow(env):
    v = H("v1"); i = c.post("/api/v1/issues", json={"title": "Mic dead", "description": "No audio", "event_id": env["e1"], "severity": "HIGH"}, headers=v).json()
    iid = i["id"]
    assert c.post(f"/api/v1/issues/{iid}/close", headers=v).status_code in (403, 409)
    assert c.post(f"/api/v1/issues/{iid}/transition", json={"to": "CLOSED"}, headers=H("co")).status_code == 409  # skipping states
    assert c.post(f"/api/v1/issues/{iid}/assign", json={"owner_id": env["v2"]}, headers=H("co")).json()["status"] == "ASSIGNED"
    assert c.post(f"/api/v1/issues/{iid}/transition", json={"to": "IN_PROGRESS"}, headers=H("v2")).status_code == 200
    assert c.post(f"/api/v1/issues/{iid}/resolve", json={}, headers=H("v2")).status_code == 422  # note required
    assert c.post(f"/api/v1/issues/{iid}/resolve", json={"note": "Replaced mic"}, headers=H("v2")).status_code == 200
    assert c.post(f"/api/v1/issues/{iid}/verify", headers=H("v2")).status_code == 403  # owner may not verify own work
    assert c.post(f"/api/v1/issues/{iid}/verify", headers=H("co")).status_code == 200
    assert c.post(f"/api/v1/issues/{iid}/close", headers=H("co")).json()["status"] == "CLOSED"
    assert c.get(f"/api/v1/issues/{iid}", headers=H("co")).status_code == 200  # closed issues stay visible
    assert len(c.get(f"/api/v1/issues/{iid}/history", headers=H("co")).json()) >= 6

def test_escalation_is_idempotent(env):
    db = SessionLocal(); i = Issue(ref="FO-OLD", title="old", description="d", event_id=env["e1"], category="x", severity="CRITICAL",
        respond_by=now() - dt.timedelta(hours=3), resolve_by=now() - dt.timedelta(hours=1)); db.add(i); db.commit()
    cron = {"Authorization": "Bearer cron-secret"}
    assert c.post("/api/v1/jobs/escalate").status_code == 401
    first = c.post("/api/v1/jobs/escalate", headers=cron).json()["escalations_created"]
    assert first >= 2 and c.post("/api/v1/jobs/escalate", headers=cron).json()["escalations_created"] == 0

def test_dashboard_counts_records(env):
    s = c.get("/api/v1/dashboard/summary", headers=H("core")).json()
    assert s["duties_scheduled"] == 2 and s["overdue_issues"] >= 1

def test_core_creates_member_and_changes_role(env):
    body = {"email": "New.Person@t.invalid", "name": "New Person", "role": "VOLUNTEER", "password": "a-long-password-1"}
    r = c.post("/api/v1/users", json=body, headers=H("core")); assert r.status_code == 201 and r.json()["email"] == "new.person@t.invalid" and "pw_hash" not in r.json()
    uid = r.json()["id"]
    assert c.post("/api/v1/auth/login", json={"email": "new.person@t.invalid", "password": "a-long-password-1"}).status_code == 200
    assert c.post(f"/api/v1/users/{uid}/roles", json={"role": "COORDINATOR"}, headers=H("dep")).status_code == 403
    assert c.post(f"/api/v1/users/{uid}/roles", json={"role": "COORDINATOR"}, headers=H("core")).json()["role"] == "COORDINATOR"
    assert c.patch(f"/api/v1/users/{uid}", json={"phone": "9000000000"}, headers=H("v1")).status_code == 403
    assert c.post(f"/api/v1/users/{uid}/deactivate", headers=H("core")).status_code == 200
    assert c.post("/api/v1/auth/login", json={"email": "new.person@t.invalid", "password": "a-long-password-1"}).status_code == 401

def test_login_throttle(env):
    for _ in range(5): c.post("/api/v1/auth/login", json={"email": "v2@t.invalid", "password": "nope-nope-nope"})
    assert c.post("/api/v1/auth/login", json={"email": "v2@t.invalid", "password": PW}).status_code == 429

def test_reassign_keeps_history_and_calendar(env):
    d = c.get("/api/v1/duties", params={"session_id": env["s1"]}, headers=H("core")).json()[0]
    new = c.post("/api/v1/users", json={"email": "cover@t.invalid", "name": "Cover", "role": "VOLUNTEER", "password": PW + "x"}, headers=H("core")).json()["id"]
    r = c.post(f"/api/v1/duties/{d['id']}/reassign", json={"user_id": new, "reason": "Illness"}, headers=H("core")); assert r.status_code == 201 or r.status_code == 200
    assert c.post(f"/api/v1/duties/{d['id']}/reassign", json={"user_id": new, "reason": "again"}, headers=H("core")).status_code == 409  # old duty is now cancelled
    assert c.post(f"/api/v1/duties/{d['id']}/cancel", json={"reason": "x"}, headers=H("core")).status_code == 422  # reason too short
    start = (now() - dt.timedelta(days=1)).isoformat(); end = (now() + dt.timedelta(days=1)).isoformat()
    cal = c.get("/api/v1/calendar", params={"start": start, "end": end}, headers=H("core")).json()
    s1 = next(s for s in cal if s["id"] == env["s1"]); assert [x["user_id"] for x in s1["staff"]] == [new] and "POC" not in s1["missing"]
    assert all(s["event_id"] == env["e1"] for s in c.get("/api/v1/calendar", params={"start": start, "end": end}, headers=H("co")).json())

def make_xlsx():
    import io
    from openpyxl import Workbook
    wb = Workbook(); m = wb.active; m.title = "Member Directory"; m.append(["#", "Name", "Role", "x", "Reports", "Phone"]); m.append([1, "Imported Person", "Coordinator", None, "-", 9876543210])
    r = wb.create_sheet("Online Duty Roster"); r.append(["S.No.", "Event ID", "Event Name", "Stage", "Date", "Start", "End", "Sports POC", "Backup Support", "Coordinator"])
    r.append([1, "IMP-1", "Imported Event", "Final", 46289, 0.5, 0.6, "Imported Person", "-", "-"]); r.append([None] * 10)  # formatted empty row must be ignored
    i = wb.create_sheet("Issue Log"); i.append(["ID", "Date", "Time", "Event", "Stage", "Sev", "Issue", "Owner", "Status", "Resolution", "Closed"]); i.append(["ISS-1", "24-09-2026", 0.4, "Unknown Event", "F", "HIGH", "Mic", "X", "Open", None, None])
    b = io.BytesIO(); wb.save(b); return b.getvalue()

def test_import_requires_preview_and_is_idempotent(env):
    data, h = make_xlsx(), H("core"); P = "/api/v1/imports/workbook"
    assert c.post(P + "/preview", content=data, headers=H("dep")).status_code == 403
    r = c.post(P + "/preview", content=data, headers=h).json(); rep = r["report"]
    assert rep["new"]["events"] == 1 and rep["new"]["duties"] == 1 and rep["new"]["issues"] == 0 and len(rep["errors"]) == 1  # unknown event reported, not invented
    assert SessionLocal().query(Event).filter_by(code="IMP-1").count() == 0  # preview saved nothing
    assert c.post(P + "/commit?sha256=" + "0" * 64, content=data, headers=h).status_code == 409
    assert c.post(P + f"/commit?sha256={r['sha256']}", content=data, headers=h).status_code == 200
    assert SessionLocal().query(Event).filter_by(code="IMP-1").count() == 1
    again = c.post(P + f"/commit?sha256={r['sha256']}", content=data, headers=h).json()["report"]
    assert again["new"]["events"] == 0 and again["new"]["duties"] == 0 and again["existing"]["duties"] == 1
    assert "Unknown Event" in c.get(f"/api/v1/imports/{r['job_id']}/errors", headers=h).text

def test_attachments_validated_and_access_controlled(env):
    i = c.post("/api/v1/issues", json={"title": "=HYPERLINK(1)", "description": "d", "event_id": env["e1"], "severity": "LOW"}, headers=H("v1")).json()["id"]; U = f"/api/v1/issues/{i}/attachments"
    png = b"\x89PNG\r\n\x1a\n" + b"0" * 32
    a = c.post(U + "?filename=../../evil.png", content=png, headers=H("v1")); assert a.status_code == 201 and "/" not in a.json()["filename"] and a.json()["content_type"] == "image/png"
    assert c.post(U + "?filename=x.png", content=b"<html>not an image</html>", headers=H("v1")).status_code == 415
    assert c.post(U + "?filename=x.jpg", content=b"\xff\xd8\xff" + b"0" * (3 * 1024 * 1024), headers=H("v1")).status_code == 413
    assert c.get(f"/api/v1/attachments/{a.json()['id']}", headers=H("v1")).content == png
    assert c.get(f"/api/v1/attachments/{a.json()['id']}", headers=H("v2")).status_code == 404  # another volunteer cannot read it
    assert c.get(U, headers=H("co")).status_code == 200

def test_exports_scoped_and_formula_safe(env):
    r = c.get("/api/v1/exports/issues?format=csv", headers=H("core")); assert r.status_code == 200 and "Reference" in r.text and "'=HYPERLINK(1)" in r.text
    assert c.get("/api/v1/exports/issues?format=xlsx", headers=H("core")).content[:2] == b"PK"
    assert c.get("/api/v1/exports/attendance", headers=H("core")).status_code == 200
    assert c.get("/api/v1/exports/issues", headers=H("v1")).status_code == 403
    scoped = c.get("/api/v1/exports/duties", headers=H("co")).text; assert "Imported Event" not in scoped  # coordinator scope excludes unrelated events

def test_audit_log_listing(env):
    d = c.get("/api/v1/audit-logs?action=duty", headers=H("dep")).json()
    assert d["total"] >= 3 and all(x["action"].startswith("duty") for x in d["items"]) and d["items"][0]["actor_name"]
    assert c.get("/api/v1/audit-logs", headers=H("co")).status_code == 403
